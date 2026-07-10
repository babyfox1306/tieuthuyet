"""Catalog — promote, migrate, export, morning report."""

from __future__ import annotations

import json
import random
import re
import shutil
import unicodedata
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import yaml

from factory.engine.lib.language import target_language
from factory.engine.lib.machine_qc import format_machine_reasons, issues_to_needs_fix, machine_qc, word_count_vi
from factory.engine.lib.prompt_builder import load_direction
from factory.engine.lib.qc_eval import format_qc_reasons
from factory.engine.lib.epub_qc import format_epub_qc_summary, qc_epub_file
from factory.engine.lib.export_gate import (
    assert_export_gate,
    check_chapter_for_promote,
    export_gate_pass,
    format_export_gate_reasons,
    prepare_chapter_for_export,
)
from factory.engine.lib.prose_sanitize import sanitize_prose
from factory.engine.paths import (
    CATALOG,
    book_catalog_dir,
    book_workspace_dir,
    catalog_archive_dir,
    catalog_series_dir,
    chapter_pipeline_path,
    load_config,
    pipeline_dir,
    promoted_marker,
    resolve_book_slug,
    workspace_dir,
)

def slugify_vi(text: str) -> str:
    s = unicodedata.normalize("NFD", text)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D").lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def is_generic_chapter_title(title: str, chapter: int, lang: str = "en") -> bool:
    """True when title is only 'Chapter 16' / 'Chương 16' with no real name."""
    t = title.strip()
    if not t:
        return True
    if lang == "en":
        return bool(re.fullmatch(rf"Chapter\s+{chapter}\s*", t, re.IGNORECASE))
    return bool(re.fullmatch(rf"Chương\s+{chapter}\s*", t, re.IGNORECASE))


def chapter_beat_from_plan(workspace_id: str, book: int, chapter: int) -> dict | None:
    """Beat dict from outline / master_plan for a chapter number."""
    ws = workspace_dir(workspace_id)
    book_dir = ws / "books" / f"{book:02d}"
    for name in ("outline.json", "master_plan.json"):
        path = book_dir / name
        if not path.exists():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        for key in ("chapter_plans", "chapter_beats", "chapters", "beats"):
            beats = data.get(key)
            if not isinstance(beats, list):
                continue
            for beat in beats:
                if beat.get("chapter") == chapter:
                    return beat
    return None


def chapter_title_from_plan(workspace_id: str, book: int, chapter: int) -> str | None:
    """Canonical chapter title from outline / master_plan (source of truth when prose has no # heading)."""
    beat = chapter_beat_from_plan(workspace_id, book, chapter)
    if not beat:
        return None
    t = beat.get("title")
    return str(t).strip() if t and str(t).strip() else None


def resolve_chapter_display(
    workspace_id: str,
    meta: dict,
    body: str,
    lang: str,
) -> tuple[str, str | None, str]:
    """Title/subtitle/body for export — never emit bare 'Chapter N' if plan has a real title."""
    title, subtitle, body = prepare_chapter_for_export(meta, body, lang)
    ch = int(meta.get("chapter") or 0)
    book = int(meta.get("book") or 1)
    if is_generic_chapter_title(title, ch, lang):
        plan_title = chapter_title_from_plan(workspace_id, book, ch)
        if plan_title:
            title = plan_title
    return title, subtitle, body


def parse_chapter_header(text: str) -> tuple[int | None, str | None, str]:
    lines = text.strip().splitlines()
    if lines and lines[0].startswith("#"):
        m = re.match(
            r"#\s*(?:Chapter|Chương)\s*(\d+)\s*:?\s*(.*)$",
            lines[0].strip(),
            re.IGNORECASE,
        )
        if m:
            num = int(m.group(1))
            rest = m.group(2).strip()
            return num, rest or None, "\n".join(lines[1:]).strip()
    return None, None, text.strip()


def extract_subtitle(body: str) -> str:
    for line in body.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith('"') and '"' in line[1:]:
            end = line.rfind('"')
            if end > 0:
                return line[1:end]
        if not line.startswith("*") and len(line) < 120:
            return line
    return ""


def build_frontmatter(meta: dict) -> str:
    clean = {k: v for k, v in meta.items() if v is not None and v != ""}
    return "---\n" + yaml.dump(clean, allow_unicode=True, default_flow_style=False, sort_keys=False) + "---"


def normalize_catalog_chapter(
    meta: dict,
    body: str,
    lang: str,
    *,
    workspace_id: str | None = None,
    book: int | None = None,
) -> tuple[dict, str]:
    """Clean catalog source — strip markdown heading, fix EN title/subtitle."""
    from factory.engine.lib.export_gate import normalize_heading_text, strip_leading_chapter_heading

    subtitle_raw = str(meta.get("subtitle") or "")
    sub_norm = normalize_heading_text(subtitle_raw)
    chapter = int(meta.get("chapter") or 0)
    book_num = int(book if book is not None else meta.get("book") or 1)
    plan_title = (
        chapter_title_from_plan(workspace_id, book_num, chapter) if workspace_id else None
    )
    heading_title: str | None = None

    m = re.match(r"^(?:Chapter|Chương)\s+\d+\s*:\s*(.+)$", sub_norm, re.IGNORECASE)
    if m:
        heading_title = m.group(1).strip()

    first = body.lstrip("\n").split("\n", 1)[0].strip() if body.strip() else ""
    if not heading_title:
        m2 = re.match(r"^#\s*(?:Chapter|Chương)\s+\d+\s*:\s*(.+)$", first, re.IGNORECASE)
        if m2:
            heading_title = m2.group(1).strip()

    body = strip_leading_chapter_heading(body)
    body = sanitize_prose(body)
    new_meta = dict(meta)
    existing = str(meta.get("title") or "").strip()

    if lang == "en":
        if heading_title:
            new_meta["title"] = heading_title
        elif plan_title:
            new_meta["title"] = plan_title
        elif existing and not is_generic_chapter_title(existing, chapter, lang):
            new_meta["title"] = existing
        else:
            new_meta["title"] = plan_title or f"Chapter {chapter}"
        new_meta.pop("subtitle", None)
    elif heading_title and not subtitle_raw:
        new_meta["subtitle"] = heading_title

    new_meta["word_count"] = word_count_vi(body)
    return new_meta, body


def parse_markdown(path: Path) -> tuple[dict, str]:
    raw = path.read_text(encoding="utf-8")
    if not raw.startswith("---"):
        return {}, raw
    parts = raw.split("---", 2)
    if len(parts) < 3:
        return {}, raw
    meta = yaml.safe_load(parts[1]) or {}
    body = parts[2].lstrip("\n")
    return meta, body


def find_catalog_chapter_by_number(
    workspace_id: str,
    book_slug: str,
    chapter: int,
) -> Path | None:
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    if not chapters_dir.exists():
        return None
    for path in chapters_dir.glob("*.md"):
        meta, _ = parse_markdown(path)
        if int(meta.get("chapter") or 0) == chapter:
            return path
    return None


def catalog_chapter_is_clean(meta: dict) -> bool:
    """True when catalog frontmatter has no blocking needs_fix flags."""
    return not (meta.get("needs_fix") or [])


def sync_chapter_pipeline_from_catalog(
    workspace_id: str,
    book: int,
    ch: int,
    *,
    book_slug: str | None = None,
) -> bool:
    """Align pipeline/ready with clean catalog — clears stale needs_fix after manual repair."""
    cfg = load_config()
    slug = book_slug or resolve_book_slug(workspace_id, book, cfg=cfg)
    path = find_catalog_chapter_by_number(workspace_id, slug, ch)
    if not path:
        return False
    meta, body = parse_markdown(path)
    if not catalog_chapter_is_clean(meta):
        return False

    ws = workspace_dir(workspace_id)
    chapter_pipeline_path(ws, book, "ready", ch).write_text(body.strip() + "\n", encoding="utf-8")
    promoted_marker(ws, book, ch).write_text(
        datetime.now(timezone.utc).isoformat(),
        encoding="utf-8",
    )

    bd = book_workspace_dir(ws, book) / "pipeline"
    for bucket in ("needs_fix", "needs_review", "draft"):
        for pattern in (f"ch_{ch:03d}.txt", f"ch_{ch:03d}_issues.json", f"ch_{ch:03d}_qc.json"):
            stale = bd / bucket / pattern
            if stale.exists():
                stale.unlink()
    return True


def sync_pipeline_from_catalog(
    workspace_id: str,
    book: int = 1,
    *,
    book_slug: str | None = None,
) -> dict[str, Any]:
    """Sync all clean catalog chapters → pipeline ready (fixes UI/pipeline drift)."""
    cfg = load_config()
    slug = book_slug or resolve_book_slug(workspace_id, book, cfg=cfg)
    chapters_dir = book_catalog_dir(workspace_id, slug) / "chapters"
    synced: list[int] = []
    skipped: list[int] = []
    if not chapters_dir.exists():
        return {"synced": synced, "skipped": skipped}
    for path in sorted(chapters_dir.glob("*.md")):
        meta, _ = parse_markdown(path)
        ch = int(meta.get("chapter") or 0)
        if not ch:
            continue
        if sync_chapter_pipeline_from_catalog(workspace_id, book, ch, book_slug=slug):
            synced.append(ch)
        else:
            skipped.append(ch)
    return {"synced": synced, "skipped": skipped}


def chapter_filename(chapter: int, slug: str) -> str:
    return f"{chapter:02d}-{slug}.md"


def load_manifest(ws: Path) -> dict:
    path = ws / "manifest.yaml"
    if path.exists():
        return yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return {"id": ws.name}


def sync_series_yaml(workspace_id: str) -> Path:
    ws = workspace_dir(workspace_id)
    manifest = load_manifest(ws)
    series_dir = catalog_series_dir(workspace_id)
    series_dir.mkdir(parents=True, exist_ok=True)
    out = {
        "id": manifest.get("id", workspace_id),
        "pen_name": manifest.get("pen_name", ""),
        "blurb": manifest.get("blurb", "").strip(),
        "tropes": manifest.get("tropes", []),
        "spice_badge": manifest.get("spice_badge", ""),
        "target_platforms": manifest.get("target_platforms", []),
    }
    path = series_dir / "series.yaml"
    path.write_text(yaml.dump(out, allow_unicode=True, default_flow_style=False, sort_keys=False), encoding="utf-8")
    return path


def ensure_book_yaml(workspace_id: str, book_slug: str, *, book: int | None = None) -> Path:
    book_dir = book_catalog_dir(workspace_id, book_slug)
    book_dir.mkdir(parents=True, exist_ok=True)
    (book_dir / "chapters").mkdir(exist_ok=True)
    (book_dir / "exports").mkdir(exist_ok=True)
    path = book_dir / "book.yaml"
    if not path.exists():
        from factory.engine.lib.narrative_schema import load_concept

        ws = workspace_dir(workspace_id)
        concept = load_concept(ws)
        direction = load_direction(ws)
        lang = direction.get("target_language", "en")
        book_num = book or int(book_slug.split("-", 1)[0]) if book_slug[:2].isdigit() else 1
        default_title = str(concept.get("title") or book_slug)
        if book_num == 2:
            default_title = "The Altered Name"
        data = {
            "book": book_num,
            "slug": book_slug,
            "title": default_title,
            "language": lang,
            "keywords_kdp": ["contract marriage", "romance thriller", "conspiracy"],
            "categories_kdp": ["Fiction > Romance > Suspense"],
            "kindle_unlimited": True,
        }
        path.write_text(yaml.dump(data, allow_unicode=True, default_flow_style=False), encoding="utf-8")
    return path


def text_to_catalog_md(
    text: str,
    *,
    series_id: str,
    book: int,
    chapter: int,
    title: str,
    slug: str,
    spice: int = 1,
    needs_fix: list[str] | None = None,
    status: str = "draft",
    subtitle: str | None = None,
    target_lang: str | None = None,
) -> tuple[str, dict]:
    ch_num, parsed_title, body = parse_chapter_header(text)
    chapter = ch_num or chapter
    lang = (target_lang or "vi").lower()

    if parsed_title and lang == "en":
        title = parsed_title
        sub = None
    elif parsed_title:
        title = title or f"Chương {chapter}"
        sub = subtitle or parsed_title
    else:
        sub = subtitle or extract_subtitle(body)

    if is_generic_chapter_title(title, chapter, lang):
        plan_title = chapter_title_from_plan(series_id, book, chapter)
        if plan_title:
            title = plan_title

    body = sanitize_prose(body)
    wc = word_count_vi(body)
    meta = {
        "series": series_id,
        "book": book,
        "chapter": chapter,
        "title": title,
        "spice": spice,
        "word_count": wc,
        "status": status,
        "needs_fix": needs_fix or [],
        "promoted_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    if sub:
        meta["subtitle"] = sub
    return build_frontmatter(meta) + "\n\n" + body, meta


def write_catalog_chapter(
    workspace_id: str,
    book_slug: str,
    filename: str,
    content: str,
    *,
    spot_check: bool = True,
) -> Path:
    series_dir = catalog_series_dir(workspace_id)
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    chapters_dir.mkdir(parents=True, exist_ok=True)
    out = chapters_dir / filename
    out.write_text(content, encoding="utf-8")
    if spot_check:
        sc_dir = series_dir / "spot_check"
        sc_dir.mkdir(exist_ok=True)
        shutil.copy2(out, sc_dir / filename)
    return out


def backup_catalog_book(workspace_id: str, book_slug: str) -> Path:
    """Rolling backup: one archive folder per book; each run replaces the last."""
    src = book_catalog_dir(workspace_id, book_slug)
    if not src.exists():
        raise FileNotFoundError(f"catalog not found: {src}")
    dest = catalog_archive_dir(workspace_id, book_slug)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists():
        shutil.rmtree(dest)
    shutil.copytree(src, dest)
    meta = {
        "workspace_id": workspace_id,
        "book_slug": book_slug,
        "backed_up_at": datetime.now(timezone.utc).isoformat(),
    }
    (dest / ".backup_meta.json").write_text(
        json.dumps(meta, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return dest


def maybe_auto_backup_catalog(
    workspace_id: str,
    book_slug: str,
    cfg: dict | None = None,
    *,
    reason: str = "export",
) -> Path | None:
    """Rolling backup before destructive catalog writes when enabled in config."""
    cfg = cfg or load_config()
    key = "auto_backup_on_export" if reason == "export" else "auto_backup_on_promote"
    if not cfg.get(key, True):
        return None
    src = book_catalog_dir(workspace_id, book_slug)
    if not src.exists() or not (src / "chapters").exists():
        return None
    dest = backup_catalog_book(workspace_id, book_slug)
    safe_print(f"[backup] {reason} -> {dest}")
    return dest


def repair_catalog_book(
    workspace_id: str,
    book_slug: str,
    *,
    backup: bool = True,
) -> dict[str, Any]:
    """Normalize catalog chapters, refresh machine_qc needs_fix, optional backup."""
    cfg = load_config()
    ws = workspace_dir(workspace_id)
    direction = load_direction(ws)
    backup_path: Path | None = None
    if backup:
        backup_path = maybe_auto_backup_catalog(workspace_id, book_slug, cfg, reason="repair")

    lang = export_language(workspace_id)
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    fixed = 0
    cleared_flags = 0
    for path in sorted(chapters_dir.glob("*.md")):
        meta, body = parse_markdown(path)
        book_num = int(meta.get("book") or 1)
        m_issues = machine_qc(
            body,
            min_words=cfg["min_word_count"],
            direction=direction,
            cfg=cfg,
            target_lang=lang,
            workspace_id=workspace_id,
            book=book_num,
        )
        needs_fix = issues_to_needs_fix(m_issues)
        if meta.get("needs_fix") and not needs_fix:
            cleared_flags += 1
        meta["needs_fix"] = needs_fix
        meta["word_count"] = word_count_vi(sanitize_prose(body))
        new_meta, new_body = normalize_catalog_chapter(
            meta, body, lang, workspace_id=workspace_id, book=book_num
        )
        content = build_frontmatter(new_meta) + "\n\n" + new_body
        write_catalog_chapter(workspace_id, book_slug, path.name, content, spot_check=True)
        fixed += 1
        if catalog_chapter_is_clean(new_meta):
            sync_chapter_pipeline_from_catalog(
                workspace_id, book_num, int(new_meta.get("chapter") or 0), book_slug=book_slug
            )

    return {
        "fixed": fixed,
        "cleared_needs_fix": cleared_flags,
        "backup": str(backup_path) if backup_path else None,
        "target_language": lang,
    }


def promote_chapter(
    workspace_id: str,
    book: int,
    ch: int,
    *,
    book_slug: str | None = None,
    spice: int | None = None,
    auto: bool = False,
) -> tuple[Path | None, list[str]]:
    cfg = load_config()
    ws = workspace_dir(workspace_id)
    series_id = load_manifest(ws).get("id", workspace_id)
    book_slug = book_slug or resolve_book_slug(workspace_id, book, cfg=cfg)

    ready_path = ws / "books" / f"{book:02d}" / "pipeline" / "ready" / f"ch_{ch:03d}.txt"
    if not ready_path.exists():
        return None, []
    if promoted_marker(ws, book, ch).exists():
        return None, []

    text = ready_path.read_text(encoding="utf-8")
    ch_num, parsed_title, _ = parse_chapter_header(text)
    ch_num = ch_num or ch
    lang = export_language(workspace_id)
    plan_beat = chapter_beat_from_plan(workspace_id, book, ch_num)

    if parsed_title:
        title = parsed_title
    elif plan_beat and plan_beat.get("title"):
        title = str(plan_beat["title"]).strip()
    else:
        title = f"Chapter {ch_num}" if lang == "en" else f"Chương {ch_num}"

    if plan_beat and plan_beat.get("slug"):
        slug = str(plan_beat["slug"]).strip()
    else:
        slug = slugify_vi(title)

    if spice is None and plan_beat and plan_beat.get("spice") is not None:
        spice = int(plan_beat["spice"])

    spice = spice if spice is not None else (3 if ch_num == 3 else 1)

    direction = load_direction(ws)
    m_issues = machine_qc(
        text,
        min_words=cfg["min_word_count"],
        direction=direction,
        cfg=cfg,
        workspace_id=workspace_id,
        book=book,
    )
    needs_fix = issues_to_needs_fix(m_issues)

    md, meta = text_to_catalog_md(
        text,
        series_id=series_id,
        book=book,
        chapter=ch_num,
        title=title,
        slug=slug,
        spice=spice,
        needs_fix=needs_fix,
        target_lang=lang,
    )
    body_for_gate = md.split("---", 2)[-1].lstrip("\n") if md.startswith("---") else md
    lang = export_language(workspace_id)
    gate_checks = check_chapter_for_promote(
        meta, body_for_gate, lang, cfg=cfg, workspace_id=workspace_id
    )
    if gate_checks and not export_gate_pass({"checks": gate_checks}, cfg):
        reasons = format_export_gate_reasons({"checks": gate_checks})
        safe_print(f"  [promote blocked] ch_{ch:03d}: export gate")
        for reason in reasons[:6]:
            safe_print(f"    {reason}")
        return None, reasons

    from factory.engine.lib.canon_guard import format_canon_guard_reasons, run_canon_guard

    cg = run_canon_guard(workspace_id, body_for_gate, chapter=ch_num, book=book)
    if not cg.get("passed"):
        reasons = format_canon_guard_reasons(cg)
        safe_print(f"  [promote blocked] ch_{ch:03d}: canon guard")
        for reason in reasons[:6]:
            safe_print(f"    {reason}")
        return None, reasons

    fname = chapter_filename(ch_num, slug)
    existing = find_catalog_chapter_by_number(workspace_id, book_slug, ch_num)
    if existing:
        maybe_auto_backup_catalog(workspace_id, book_slug, cfg, reason="promote")
        fname = existing.name
    elif promoted_marker(ws, book, ch).exists():
        return None, []

    sync_series_yaml(workspace_id)
    ensure_book_yaml(workspace_id, book_slug, book=book)
    out = write_catalog_chapter(workspace_id, book_slug, fname, md)
    promoted_marker(ws, book, ch).write_text(datetime.now(timezone.utc).isoformat(), encoding="utf-8")
    tag = "re-promoted" if existing else ("auto-promoted" if auto else "promoted")
    print(f"  [{tag}] ch_{ch:03d} -> catalog/{series_id}/.../{fname}")
    return out, []


def promote_all(workspace_id: str, book: int = 1) -> int:
    ws = workspace_dir(workspace_id)
    ready_dir = ws / "books" / f"{book:02d}" / "pipeline" / "ready"
    if not ready_dir.exists():
        return 0
    count = 0
    for path in sorted(ready_dir.glob("ch_*.txt")):
        ch = int(path.stem.split("_")[1])
        out, _reasons = promote_chapter(workspace_id, book, ch, auto=False)
        if out:
            count += 1
    return count


MIGRATE_MAP = [
    {
        "file": "chapter1_output.txt",
        "chapter": 1,
        "title": "Hợp đồng có giá",
        "slug": "hop-dong-co-gia",
        "subtitle": "Một tỷ đồng. Ký tên vào đây.",
        "spice": 1,
    },
    {
        "file": "chapter2_output.txt",
        "chapter": 2,
        "title": "Mẹ Chồng Cao Tay",
        "slug": "me-chong-cao-tay",
        "subtitle": "Gọi mẹ đi con.",
        "spice": 1,
        "continuity_flags": ["continuity:ch2 mở cafe thay vì nối sảnh Lâm ch1"],
    },
    {
        "file": "chapter3_output_18plus.txt",
        "chapter": 3,
        "title": "Thử Xem",
        "slug": "thu-xem",
        "subtitle": "Tiếng gõ cửa lúc 23h47 không phải yêu cầu.",
        "spice": 3,
    },
]


def safe_print(msg: str) -> None:
    try:
        print(msg)
    except UnicodeEncodeError:
        print(msg.encode("ascii", errors="replace").decode("ascii"))


def migrate_from_scripts(
    from_dir: Path,
    workspace_id: str = "ceo-contract",
    book_slug: str = "01-hop-dong-co-gia",
    with_fixes: bool = True,
) -> int:
    cfg = load_config()
    ws = workspace_dir(workspace_id)
    series_id = load_manifest(ws).get("id", workspace_id)
    sync_series_yaml(workspace_id)
    ensure_book_yaml(workspace_id, book_slug)

    count = 0
    for spec in MIGRATE_MAP:
        src = from_dir / spec["file"]
        if not src.exists():
            print(f"  SKIP missing {src.name}")
            continue
        text = src.read_text(encoding="utf-8")
        extra: list[str] = []
        if with_fixes:
            direction = load_direction(ws)
            m_issues = machine_qc(
                text,
                min_words=cfg["min_word_count"],
                direction=direction,
                cfg=cfg,
                workspace_id=workspace_id,
                book=1,
            )
            extra = issues_to_needs_fix(m_issues)
            extra.extend(spec.get("continuity_flags", []))

        md, meta = text_to_catalog_md(
            text,
            series_id=series_id,
            book=1,
            chapter=spec["chapter"],
            title=spec["title"],
            slug=spec["slug"],
            spice=spec["spice"],
            needs_fix=extra,
            subtitle=spec.get("subtitle"),
        )
        fname = chapter_filename(spec["chapter"], spec["slug"])
        write_catalog_chapter(workspace_id, book_slug, fname, md, spot_check=True)
        safe_print(f"  migrated {src.name} -> {fname} needs_fix={meta['needs_fix']}")
        count += 1

    return count


def reader_title(meta: dict) -> str:
    title = meta.get("title", "")
    sub = meta.get("subtitle", "")
    if sub:
        return f"Chương {meta.get('chapter', '')} — {title}\n{sub}"
    return f"Chương {meta.get('chapter', '')} — {title}"


EPUB_CSS = """body {
  margin: 1em 5%;
  font-family: Georgia, "Times New Roman", serif;
  line-height: 1.5;
}
h1 {
  font-size: 1.4em;
  margin: 1.2em 0 0.6em;
  page-break-before: always;
}
h1:first-of-type {
  page-break-before: auto;
}
p {
  margin: 0.6em 0;
  text-indent: 0;
}
em.subtitle {
  font-style: italic;
  display: block;
  margin-bottom: 1em;
}
"""


def _escape_xml(text: str) -> str:
    return escape(str(text), {"'": "&apos;", '"': "&quot;"})


def load_chapter_items(chapters_dir: Path) -> list[tuple[Path, dict, str]]:
    items: list[tuple[Path, dict, str]] = []
    for path in chapters_dir.glob("*.md"):
        meta, body = parse_markdown(path)
        items.append((path, meta, body))
    items.sort(key=lambda x: (int(x[1].get("chapter", 0) or 0), x[0].name))
    return items


def _load_book_yaml(workspace_id: str, book_slug: str) -> dict:
    path = book_catalog_dir(workspace_id, book_slug) / "book.yaml"
    if not path.exists():
        return {}
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


def _save_book_yaml(workspace_id: str, book_slug: str, data: dict) -> None:
    path = book_catalog_dir(workspace_id, book_slug) / "book.yaml"
    path.write_text(yaml.dump(data, allow_unicode=True, default_flow_style=False, sort_keys=False), encoding="utf-8")


def export_language(workspace_id: str) -> str:
    ws = workspace_dir(workspace_id)
    direction = load_direction(ws)
    cfg = load_config()
    lang = target_language(direction, cfg)
    bible = ws / "bible" / "series.json"
    if bible.exists():
        try:
            data = json.loads(bible.read_text(encoding="utf-8"))
            lang = data.get("meta", {}).get("target_language", lang)
        except json.JSONDecodeError:
            pass
    return str(lang).strip().lower().split("-")[0] or "vi"


def ensure_epub_identifier(workspace_id: str, book_slug: str) -> str:
    data = _load_book_yaml(workspace_id, book_slug)
    uid = data.get("epub_identifier")
    if not uid:
        uid = f"urn:uuid:{uuid.uuid4()}"
        data["epub_identifier"] = uid
        _save_book_yaml(workspace_id, book_slug, data)
    return uid


def book_display_title(workspace_id: str, book_slug: str) -> str:
    data = _load_book_yaml(workspace_id, book_slug)
    return str(data.get("title") or book_slug)


def _zip_epub_write(
    zf: zipfile.ZipFile,
    name: str,
    data: bytes | str,
    *,
    compress: int,
) -> None:
    info = zipfile.ZipInfo(name)
    info.compress_type = compress
    payload = data.encode("utf-8") if isinstance(data, str) else data
    zf.writestr(info, payload)


def _cover_media_type(cover_path: Path) -> tuple[str, str]:
    ext = cover_path.suffix.lower()
    if ext in (".jpg", ".jpeg"):
        return "image/jpeg", "cover.jpg"
    if ext == ".png":
        return "image/png", "cover.png"
    raise ValueError(f"Cover must be .jpg or .png, got {cover_path.suffix}")


def _chapter_xhtml(title: str, subtitle: str | None, body: str, lang: str) -> str:
    parts = [
        '<?xml version="1.0" encoding="utf-8"?>',
        f'<html xmlns="http://www.w3.org/1999/xhtml" xml:lang="{lang}" lang="{lang}">',
        "<head>",
        f"<title>{_escape_xml(title)}</title>",
        '<link rel="stylesheet" type="text/css" href="style.css"/>',
        "</head>",
        "<body>",
        f"<h1>{_escape_xml(title)}</h1>",
    ]
    if subtitle:
        parts.append(f'<p class="subtitle"><em>{_escape_xml(subtitle)}</em></p>')
    for para in body.split("\n\n"):
        p = para.strip()
        if p:
            parts.append(f"<p>{_escape_xml(p)}</p>")
    parts.extend(["</body>", "</html>"])
    return "\n".join(parts)


def export_epub(
    workspace_id: str,
    book_slug: str,
    out_dir: Path,
    *,
    cover_path: Path | None = None,
) -> Path:
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    out_dir.mkdir(parents=True, exist_ok=True)
    epub_path = out_dir / f"{book_slug}.epub"
    items = load_chapter_items(chapters_dir)
    if not items:
        raise FileNotFoundError("No chapters to export")

    lang = export_language(workspace_id)
    book_title = book_display_title(workspace_id, book_slug)
    epub_uid = ensure_epub_identifier(workspace_id, book_slug)
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    chapter_entries: list[tuple[str, str, dict, str]] = []
    for i, (_path, meta, body) in enumerate(items, 1):
        title, subtitle, body = resolve_chapter_display(workspace_id, meta, body, lang)
        fname = f"chapter_{i:03d}.xhtml"
        chapter_entries.append((fname, _chapter_xhtml(title, subtitle, body, lang), meta, title))

    cover_href: str | None = None
    cover_bytes: bytes | None = None
    cover_mime: str | None = None
    if cover_path:
        cover_path = Path(cover_path)
        if not cover_path.is_file():
            raise FileNotFoundError(f"Cover not found: {cover_path}")
        cover_mime, cover_href = _cover_media_type(cover_path)
        cover_bytes = cover_path.read_bytes()

    nav_items = "".join(
        f'<li><a href="{fname}">{_escape_xml(display_title)}</a></li>'
        for fname, _html, meta, display_title in chapter_entries
    )
    nav_xhtml = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        f'<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
        f'xml:lang="{lang}" lang="{lang}">\n'
        "<head>\n"
        f"<title>{_escape_xml(book_title)}</title>\n"
        '<link rel="stylesheet" type="text/css" href="style.css"/>\n'
        "</head>\n"
        "<body>\n"
        '<nav epub:type="toc" id="toc">\n'
        f"<ol>\n{nav_items}\n</ol>\n"
        "</nav>\n"
        "</body>\n</html>"
    )

    manifest = [
        '<item id="style" href="style.css" media-type="text/css"/>',
        '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
    ]
    if cover_href and cover_mime:
        manifest.append(
            f'<item id="cover-image" href="{cover_href}" media-type="{cover_mime}" properties="cover-image"/>'
        )
    for i, (fname, _html, _meta, _title) in enumerate(chapter_entries, 1):
        manifest.append(
            f'<item id="ch{i}" href="{fname}" media-type="application/xhtml+xml"/>'
        )

    spine = "".join(f'<itemref idref="ch{i}"/>' for i in range(1, len(chapter_entries) + 1))

    meta_block = (
        f'<dc:identifier id="uid">{_escape_xml(epub_uid)}</dc:identifier>'
        f"<dc:title>{_escape_xml(book_title)}</dc:title>"
        f"<dc:language>{_escape_xml(lang)}</dc:language>"
        f'<meta property="dcterms:modified">{modified}</meta>'
    )
    if cover_href:
        meta_block += '<meta name="cover" content="cover-image"/>'

    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="uid" '
        f'xml:lang="{lang}">\n'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        f"{meta_block}\n"
        "</metadata>\n"
        "<manifest>\n"
        + "\n".join(manifest)
        + "\n</manifest>\n"
        f"<spine>\n{spine}\n</spine>\n"
        "</package>"
    )

    container_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
        "  <rootfiles>\n"
        '    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>\n'
        "  </rootfiles>\n"
        "</container>"
    )

    with zipfile.ZipFile(epub_path, "w") as zf:
        _zip_epub_write(zf, "mimetype", "application/epub+zip", compress=zipfile.ZIP_STORED)
        _zip_epub_write(zf, "META-INF/container.xml", container_xml, compress=zipfile.ZIP_DEFLATED)
        _zip_epub_write(zf, "OEBPS/content.opf", opf, compress=zipfile.ZIP_DEFLATED)
        _zip_epub_write(zf, "OEBPS/nav.xhtml", nav_xhtml, compress=zipfile.ZIP_DEFLATED)
        _zip_epub_write(zf, "OEBPS/style.css", EPUB_CSS, compress=zipfile.ZIP_DEFLATED)
        if cover_href and cover_bytes is not None:
            _zip_epub_write(zf, f"OEBPS/{cover_href}", cover_bytes, compress=zipfile.ZIP_DEFLATED)
        for fname, html, _meta, _title in chapter_entries:
            _zip_epub_write(zf, f"OEBPS/{fname}", html, compress=zipfile.ZIP_DEFLATED)

    if not cover_path:
        safe_print("[export] WARN: chưa có cover — không publish KDP được (dùng --cover path.jpg)")

    return epub_path


def export_docx(workspace_id: str, book_slug: str, out_dir: Path) -> Path:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("pip install python-docx để export --target docx") from exc

    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / f"{book_slug}.docx"
    items = load_chapter_items(chapters_dir)
    if not items:
        raise FileNotFoundError("No chapters to export")

    doc = Document()
    book_title = book_display_title(workspace_id, book_slug)
    doc.add_heading(book_title, level=0)
    lang = export_language(workspace_id)

    for _path, meta, body in items:
        title, subtitle, body = resolve_chapter_display(workspace_id, meta, body, lang)
        doc.add_heading(title, level=1)
        if subtitle:
            p = doc.add_paragraph(str(subtitle))
            p.style = "Subtitle" if "Subtitle" in [s.name for s in doc.styles] else p.style
            for run in p.runs:
                run.italic = True
        for para in body.split("\n\n"):
            text = para.strip()
            if text:
                doc.add_paragraph(text)

    doc.save(docx_path)
    return docx_path


def validate_epub_structure(epub_path: Path) -> list[str]:
    """Quick structural checks aligned with EPUB / KDP requirements."""
    issues: list[str] = []
    with zipfile.ZipFile(epub_path) as zf:
        names = zf.namelist()
        if not names or names[0] != "mimetype":
            issues.append("mimetype:not_first_entry")
        if "mimetype" in zf.namelist():
            info = zf.getinfo("mimetype")
            if info.compress_type != zipfile.ZIP_STORED:
                issues.append("mimetype:not_stored")
            if zf.read("mimetype") != b"application/epub+zip":
                issues.append("mimetype:wrong_content")
        required = (
            "META-INF/container.xml",
            "OEBPS/content.opf",
            "OEBPS/nav.xhtml",
            "OEBPS/style.css",
        )
        for req in required:
            if req not in names:
                issues.append(f"missing:{req}")
        opf = zf.read("OEBPS/content.opf").decode("utf-8")
        if 'id="uid"' not in opf or "dc:identifier" not in opf:
            issues.append("opf:missing_dc_identifier")
        if "properties=\"nav\"" not in opf:
            issues.append("opf:nav_not_marked")
        if "style.css" not in opf:
            issues.append("opf:missing_css_manifest")
        nav = zf.read("OEBPS/nav.xhtml").decode("utf-8")
        if 'epub:type="toc"' not in nav:
            issues.append("nav:missing_epub_type_toc")
        if "Ch 1" in nav or "Ch 2" in nav:
            issues.append("nav:generic_ch_labels")
        ch1 = zf.read("OEBPS/chapter_001.xhtml").decode("utf-8")
        if 'lang="' not in ch1 or "style.css" not in ch1:
            issues.append("chapter:missing_lang_or_css")
    return issues


def export_vella(workspace_id: str, book_slug: str, out_dir: Path) -> int:
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    out_dir.mkdir(parents=True, exist_ok=True)
    n = 0
    for path in sorted(chapters_dir.glob("*.md")):
        meta, body = parse_markdown(path)
        header = reader_title(meta)
        out = out_dir / f"{path.stem}.txt"
        out.write_text(f"{header}\n\n{body}", encoding="utf-8")
        n += 1
    return n


def export_kdp(workspace_id: str, book_slug: str, out_dir: Path) -> Path:
    """Deprecated — use export_docx. Kept for internal migration scripts only."""
    return export_docx(workspace_id, book_slug, out_dir)


def export_book(
    workspace_id: str,
    book_slug: str,
    target: str,
    *,
    cover_path: Path | None = None,
    cfg: dict | None = None,
) -> list[Path]:
    cfg = cfg or load_config()
    if cfg.get("export_gate_enabled", True):
        gate_report = assert_export_gate(workspace_id, book_slug, cfg=cfg)
        safe_print(f"[export] {gate_report.get('summary', 'export gate OK')}")
    else:
        safe_print("[export] export gate skipped (disabled in config)")

    maybe_auto_backup_catalog(workspace_id, book_slug, cfg, reason="export")

    exports = book_catalog_dir(workspace_id, book_slug) / "exports"
    outputs: list[Path] = []
    if target == "vella":
        n = export_vella(workspace_id, book_slug, exports / "vella")
        print(f"[export] vella: {n} chapter txt -> {exports / 'vella'}")
        outputs.append(exports / "vella")
    elif target == "docx":
        p = export_docx(workspace_id, book_slug, exports / "docx")
        print(f"[export] docx: {p}")
        outputs.append(p)
    elif target == "epub":
        p = export_epub(workspace_id, book_slug, exports / "epub", cover_path=cover_path)
        issues = validate_epub_structure(p)
        print(f"[export] epub: {p}")
        if issues:
            safe_print(f"[export] epub structure issues: {issues}")
        else:
            safe_print("[export] epub structure checks: PASS")
        qc = qc_epub_file(p)
        safe_print(f"[export] {format_epub_qc_summary(qc)}")
        if qc.get("messages"):
            for msg in qc["messages"][:10]:
                safe_print(f"  {msg['level']}({msg['code']}): {msg['message'][:200]}")
            if len(qc["messages"]) > 10:
                safe_print(f"  ... +{len(qc['messages']) - 10} more (see .qc.json)")
        if not qc.get("passed", False):
            safe_print(f"[export] WARN: EPUBCheck did not pass — report: {p.with_suffix('.qc.json')}")
        outputs.append(p)
    else:
        raise ValueError(f"Unknown target: {target} (use vella|epub|docx)")
    return outputs


def list_catalog_chapters(workspace_id: str, book_slug: str) -> list[dict]:
    chapters_dir = book_catalog_dir(workspace_id, book_slug) / "chapters"
    result = []
    for path in sorted(chapters_dir.glob("*.md")):
        meta, _ = parse_markdown(path)
        meta["file"] = path.name
        result.append(meta)
    return result


def pipeline_counts(ws: Path, book: int) -> dict[str, list[int]]:
    base = ws / "books" / f"{book:02d}" / "pipeline"
    out: dict[str, list[int]] = {}
    for bucket in ("needs_fix", "needs_review", "ready", "draft"):
        d = base / bucket
        if not d.exists():
            out[bucket] = []
            continue
        nums = []
        for p in d.glob("ch_*.txt"):
            nums.append(int(p.stem.split("_")[1]))
        out[bucket] = sorted(nums)
    return out


def recently_promoted(workspace_id: str) -> list[str]:
    sc = catalog_series_dir(workspace_id) / "spot_check"
    if not sc.exists():
        return []
    since = datetime.now(timezone.utc).timestamp() - 86400
    recent = []
    for p in sorted(sc.glob("*.md")):
        if p.stat().st_mtime >= since:
            recent.append(p.stem)
    return recent


def write_morning_report(workspace_id: str, book: int = 1, book_slug: str | None = None) -> Path:
    cfg = load_config()
    book_slug = book_slug or resolve_book_slug(workspace_id, book, cfg=cfg)
    ws = workspace_dir(workspace_id)
    manifest = load_manifest(ws)
    series_title = manifest.get("id", workspace_id).replace("-", " ").title()
    today = datetime.now().strftime("%d/%m/%Y")

    catalog = list_catalog_chapters(workspace_id, book_slug)
    ch_nums = [c["chapter"] for c in catalog]
    ch_range = f"{min(ch_nums)}–{max(ch_nums)}" if ch_nums else "—"

    pipes = pipeline_counts(ws, book)
    fix_detail = []
    for ch in pipes["needs_fix"]:
        issues_path = ws / "books" / f"{book:02d}" / "pipeline" / "needs_fix" / f"ch_{ch:03d}_issues.json"
        if issues_path.exists():
            iss = json.loads(issues_path.read_text(encoding="utf-8"))
            reasons = format_machine_reasons(iss)
            if reasons:
                fix_detail.append(f"ch{ch} ({'; '.join(reasons)})")
            else:
                keys = [k for k in ("foreign_chars", "cjk_chars", "short", "repeat") if k in iss]
                fix_detail.append(f"ch{ch} {'/'.join(keys)}")
        else:
            fix_detail.append(f"ch{ch}")

    review_detail = []
    for ch in pipes["needs_review"]:
        qc_path = ws / "books" / f"{book:02d}" / "pipeline" / "needs_review" / f"ch_{ch:03d}_qc.json"
        if qc_path.exists():
            qc = json.loads(qc_path.read_text(encoding="utf-8"))
            reasons = format_qc_reasons(qc)
            summary = "; ".join(reasons) if reasons else "QC FAIL"
            if len(summary) > 100:
                summary = summary[:97] + "…"
            review_detail.append(f"ch{ch} ({summary})")
        else:
            review_detail.append(f"ch{ch}")

    outline_total = 50
    outline_path = ws / "books" / f"{book:02d}" / "outline.json"
    if outline_path.exists():
        ol = json.loads(outline_path.read_text(encoding="utf-8"))
        outline_total = len(ol.get("chapter_beats", [])) or outline_total

    spot_pool = ch_nums or [1, 2, 3]
    spot_pick = random.sample(spot_pool, min(2, len(spot_pool)))
    recent = recently_promoted(workspace_id)
    recent_str = ", ".join(recent) if recent else "—"

    lines = [
        f"# {series_title} — {today}",
        f"✅ Catalog: {len(catalog)} chương sạch (ch {ch_range})",
        f"🆕 Auto-promoted đêm qua: {recent_str}  → đọc spot_check/",
        f"🔧 needs_fix: {len(pipes['needs_fix'])} ({', '.join(fix_detail) or '—'})",
        f"👀 needs_review: {len(pipes['needs_review'])} ({', '.join(review_detail) or '—'})",
        f"📖 Spot-check gợi ý: ch {', '.join(str(c) for c in spot_pick)}",
        f"🎯 Tiến độ book {book}: {len(catalog)}/{outline_total} chương",
    ]
    out = catalog_series_dir(workspace_id) / "MORNING.md"
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out
